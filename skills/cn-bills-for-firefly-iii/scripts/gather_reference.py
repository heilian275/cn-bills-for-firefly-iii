# -*- coding: utf-8 -*-
"""取回导入前的两侧信息：连接与 API 版本、账单列值分布（支付方式/账单分类/收支/状态）、
Firefly III 三类账户清单（资产账户附币种）、分类清单、最后 N 个月交易的
「账户/分类 ← description」对照表（锚定库内最后交易月）、本批账单逐笔明细（精简列）。

token 从 --token-file 指定的文件（每行 KEY=VALUE，找 FIREFLY_III_ACCESS_TOKEN）或环境变量读取，
不打印。

用法：
  python gather_reference.py --alipay <csv> [--wechat <xlsx>] [--url http://localhost:8080] --token-file <file> [--months 3]
"""
import argparse
import csv
import io
import sys
import urllib.error
from collections import Counter, defaultdict

import _common

AP_HEADER_ROW = 23
WX_HEADER_ROW = 17

BASELINE = "6.2"   # 脚本防护的基线版本（行为实测于此线）；实例非此线时提示复核

AP_AGG_SPEC = [
    ("交易分类", ["交易分类"]), ("收/支", ["收/支"]),
    ("支付方式", ["收/付款方式", "付款方式"]), ("状态", ["交易状态"]),
]
WX_AGG_SPEC = [
    ("交易类型", ["交易类型"]), ("收/支", ["收/支"]),
    ("支付方式", ["支付方式"]), ("状态", ["当前状态"]),
]


def col_index(header_cols, names):
    m = {c.strip(): i for i, c in enumerate(header_cols)}
    for n in names:
        for k, i in m.items():
            if n in k:
                return i
    return None


def main():
    ap = argparse.ArgumentParser(description="账单列值聚合 + Firefly III 环境与历史参照 + 当批账单明细")
    ap.add_argument("--url", default="http://localhost:8080")
    ap.add_argument("--token-file", help="含 FIREFLY_III_ACCESS_TOKEN=... 的文件路径")
    ap.add_argument("--months", type=int, default=3, help="历史参照月数（默认 3，自库内最后交易月起往前数整月）")
    ap.add_argument("--alipay", help="支付宝 csv 路径")
    ap.add_argument("--wechat", help="微信 xlsx 路径")
    args = ap.parse_args()

    token = _common.load_token(args.token_file)
    assert token, "未找到 token：请用 --token-file 或设置环境变量 FIREFLY_III_ACCESS_TOKEN"
    ff = _common.Firefly(args.url, token)

    # ---- ⓪ 连通性 + 版本（环境事实，失败即止）----
    try:
        about = ff.get("/api/v1/about")
        ver = about.get("data", {}).get("api_version") or about.get("data", {}).get("version")
        print(f"✅ 连接成功： {args.url} （API 版本 {ver}）")
        if ver and not str(ver).lstrip("v").startswith(BASELINE):
            print(f"⚠️ 实例版本非 {BASELINE}.x（脚本防护基线）；脚本异常时先按 references/firefly-iii.md 复核接口差异")
    except urllib.error.HTTPError as e:
        print(f"❌ 连接失败 HTTP {e.code}：检查 --url 与 token")
        sys.exit(1)
    except Exception as e:
        print(f"❌ 无法连接 {args.url}: {e}")
        sys.exit(1)

    # ---- 账单解析：一遍遍历，累计列值聚合 + 暂存逐笔明细 ----
    agg_blocks = []   # (标题，表头，{label: 列索引}, {label: Counter}, 行数)
    details = []      # (标题，[明细行])
    n_total = 0

    if args.alipay:
        with open(args.alipay, encoding="gb18030") as f:
            lines = f.readlines()
        hdr = [c.strip() for c in lines[AP_HEADER_ROW].rstrip().split(",")]
        ix = {n: i for i, n in enumerate(hdr)}
        idx_map = {lb: col_index(hdr, nm) for lb, nm in AP_AGG_SPEC}
        counters = {lb: Counter() for lb, _ in AP_AGG_SPEC}
        rows_out = []
        n = 0
        for ln in lines[AP_HEADER_ROW + 1:]:
            s = ln.strip()
            if not s or s.startswith("---") or ("共" in s and "笔" in s):
                continue
            row = next(csv.reader(io.StringIO(ln)))
            if len(row) < 12:
                continue
            cells = [c.strip() for c in row]
            n += 1
            n_total += 1
            for lb, _ in AP_AGG_SPEC:
                i = idx_map[lb]
                if i is not None and i < len(cells) and cells[i]:
                    counters[lb][cells[i]] += 1
            rows_out.append(f"#{n} {cells[ix['交易时间']][:16]} | {cells[ix['收/支']]} ¥{cells[ix['金额']]} | [{cells[ix['交易分类']]}] "
                            f"{cells[ix['交易对方']][:16]} - {cells[ix['商品说明']][:36]} | 付：{cells[ix['收/付款方式']] or '（空）'} | "
                            f"{cells[ix['交易状态']]} | alipay:{cells[ix['交易订单号']]}")
        agg_blocks.append(("支付宝", hdr, idx_map, counters, n))
        details.append(("支付宝账单明细", rows_out))

    if args.wechat:
        import openpyxl
        wb = openpyxl.load_workbook(args.wechat, data_only=True)
        ws = wb[wb.sheetnames[0]]
        all_rows = list(ws.iter_rows(values_only=True))
        hdr = [str(c).strip() if c is not None else "" for c in all_rows[WX_HEADER_ROW]]
        ix = {n: i for i, n in enumerate(hdr)}
        idx_map = {lb: col_index(hdr, nm) for lb, nm in WX_AGG_SPEC}
        counters = {lb: Counter() for lb, _ in WX_AGG_SPEC}
        rows_out = []
        n = 0
        for r in all_rows[WX_HEADER_ROW + 1:]:
            if r[0] is None or str(r[0]).strip() == "":
                continue
            if "共" in str(r[0]) and "笔" in str(r[0]):   # 尾部汇总行
                continue
            n += 1
            n_total += 1
            for lb, _ in WX_AGG_SPEC:
                i = idx_map[lb]
                if i is not None and r[i] is not None:
                    v = str(r[i]).strip()
                    if v:
                        counters[lb][v] += 1
            # autocorrect-disable
            rows_out.append(f"#W{n} {str(r[ix['交易时间']])[:16]} | {r[ix['收/支']]} {r[ix['金额(元)']]} | [{r[ix['交易类型']]}] "
                            f"{str(r[ix['交易对方']])[:16]} - {str(r[ix['商品']])[:36]} | 付：{r[ix['支付方式']]} | "
                            f"{r[ix['当前状态']]} | wechat:{r[ix['交易单号']]}")
            # autocorrect-enable
        agg_blocks.append(("微信", hdr, idx_map, counters, n))
        details.append(("微信账单明细", rows_out))

    # ---- ① 列值聚合 ----
    for title, hdr, idx_map, counters, n in agg_blocks:
        print("=" * 56)
        print(f"{title} 列值分布")
        print("=" * 56)
        print("表头：", hdr)
        print("列索引：", idx_map, "\n")
        for lb, c in counters.items():
            print(f"[{lb}] 共 {len(c)} 种：")
            for v, k in c.most_common():
                print(f"   {k:4d}  {v}")
            print()
        print(f"（{title}数据行 {n} 行；以上为列值聚合，逐笔明细见下方）\n")

    # ---- ①⁺ Firefly III 三类账户与分类（环境信息；币种一致性核对的数据来源）----
    asset_cur = Counter()
    try:
        for t, label in [("asset", "资产账户（支付方式/钱包）"), ("expense", "支出账户"), ("revenue", "收入账户")]:
            names = []
            try:
                for rows in ff.paged("/api/v1/accounts", {"type": t}):
                    for a in rows:
                        names.append(a["attributes"]["name"])
                        if t == "asset":
                            asset_cur[a["attributes"].get("currency_code") or "?"] += 1
            except urllib.error.HTTPError as e:
                if e.code == 404:
                    # 类型过滤端点对空类型返回 404 而非空列表（6.2.x 实测行为）
                    print(f"【{label}】0 个（空类型时此端点返回 404，非故障）")
                    continue
                raise
            print(f"【{label}】{len(names)} 个： " + "、".join(names[:30]) + ("…" if len(names) > 30 else ""))
        cur_line = "、".join(f"{c}×{n}" for c, n in asset_cur.most_common()) or "无资产账户"
        fx = [c for c in asset_cur if c not in ("RMB", "CNY")]
        print("【资产账户币种】" + cur_line + ("（全部为人民币）" if asset_cur and not fx else
              ("  ⚠️ 存在非人民币 (RMB/CNY) 账户" if fx else "")))
        cats = []
        for rows in ff.paged("/api/v1/categories"):
            cats += [a["attributes"]["name"] for a in rows]
        print(f"【分类（category）】{len(cats)} 个： " + "、".join(cats[:40]) + ("…" if len(cats) > 40 else ""))
        print()
    except urllib.error.HTTPError as e:
        print("账户/分类拉取失败 HTTP", e.code)

    # ---- ② 最后 N 个月交易 → 分类参照表（锚定库内最后交易月）----
    print("=" * 60)
    anchor = None
    try:
        anchor = ff.latest_transaction_date()
    except urllib.error.HTTPError as e:
        print("API 错误：", e.code)
    ref = defaultdict(list)
    if anchor is None:
        print("库内无交易，跳过历史参照（判读请依据上方账户/分类清单并询问用户）")
    else:
        start, end = _common.month_window(anchor, args.months)
        print(f"最后 {args.months} 个月（{start.isoformat()}～{end.isoformat()}，锚定库内最后交易月）"
              f"分类参照：方向/账户 / 分类 ← description 样本")
        print("=" * 60)
        try:
            for rows in ff.paged("/api/v1/transactions",
                                 {"start": start.isoformat(), "end": end.isoformat()}):
                for g in rows:
                    for j in g.get("attributes", {}).get("transactions", []):
                        t = j.get("type")
                        dest = j.get("destination_name") or ""
                        src = j.get("source_name") or ""
                        cat = j.get("category_name") or "（无分类）"
                        tags = ",".join(j.get("tags") or [])
                        desc = (j.get("description") or "")[:24]
                        if t == "withdrawal":
                            key = ("支出→" + dest, cat)
                        elif t == "deposit":
                            key = ("收入←" + src, cat)
                        else:
                            key = ("转账", cat)
                        ref[(key[0], cat, tags)].append(desc)
        except urllib.error.HTTPError as e:
            print("API 错误：", e.code)
        for (direction, cat, tags), descs in sorted(ref.items(), key=lambda x: -len(x[1])):
            samples = sorted(set(descs))[:4]
            print(f"  [{direction} / {cat}] x{len(descs)}: " + " | ".join(samples))
        print("\n（此表即用户既有分类体系与 description 风格的观察材料）\n")

    # ---- ③ 逐笔明细 ----
    for title, rows_out in details:
        print("=" * 60)
        print(title)
        print("=" * 60)
        for line in rows_out:
            print(line)
        print()
    print(f"（共 {n_total} 笔；token 未打印）")


if __name__ == "__main__":
    main()
